from django.test import TestCase, TransactionTestCase
from django.urls import reverse

from accounts.models import FarmProfile, User
from budgets.models import Budget

# Create your tests here.


class BudgetNavigationTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='budget-owner', password='test-password')
        self.farm = FarmProfile.objects.create(user=self.user, farm_name='Budget Test Farm')
        self.budget = Budget.objects.create(farm_profile=self.farm, name='Season Budget')
        self.client.force_login(self.user)

    def test_budget_list_links_to_registered_create_form(self):
        create_url = reverse('model_create', kwargs={'app': 'budgets', 'model': 'budget'})
        response = self.client.get('/dashboard/budgets/budget/')
        self.assertContains(response, 'New Budget')
        self.assertContains(response, reverse('budget_new'))
        self.assertEqual(self.client.get(create_url).status_code, 200)

    def test_budget_header_shows_identity_period_and_actions(self):
        self.budget.start_date = '2026-01-01'
        self.budget.end_date = '2026-12-31'
        self.budget.save()
        self.farm.location = 'Header location'
        self.farm.phone = '+256 700 123456'
        self.farm.email = 'farm@example.com'
        self.farm.save()
        response = self.client.get(reverse('budget_detail', kwargs={'pk': self.budget.pk}))
        header = response.content.decode().split('<header class="budget-toolbar"', 1)[1].split('</header>', 1)[0]
        for value in [self.farm.farm_name, self.farm.location, self.farm.phone, self.farm.email]:
            self.assertNotIn(value, header)
        self.assertNotContains(response, '<header class="letterhead">')
        for text in ['budget-header-identity', 'Annual Operating Budget',
                     'Season Budget', '01 Jan 2026', '31 Dec 2026', 'Draft',
                     'Back to Budgets', 'Edit Budget', 'Export', 'Print']:
            self.assertContains(response, text)
        self.assertContains(response, f'BDG-{self.budget.pk:05d}')

    def test_overview_is_default_and_tables_have_individual_tabs(self):
        from budgets.models import BudgetItem
        from expenses.models import Category
        category = Category.objects.create(name='Tab test code', code='TAB')
        for index, centre in enumerate(['Crops', 'Livestock'], start=1):
            BudgetItem.objects.create(budget=self.budget, category=category,
                                      cost_centre=centre, line_code=f'TAB-{index}', planned_amount=100)
        response = self.client.get(reverse('budget_detail', kwargs={'pk': self.budget.pk}))
        body = response.content.decode()
        self.assertContains(response, 'role="tablist"')
        self.assertRegex(body, r'id="tab-budget-overview"[^>]*aria-selected="true"')
        for panel in ['budget-income', 'budget-centre-1', 'budget-centre-2', 'budget-actual']:
            self.assertRegex(body, rf'id="{panel}"[^>]*hidden')
            self.assertIn(f'aria-controls="{panel}"', body)
        overview = body.split('id="budget-overview"', 1)[1].split('id="budget-table-paper"', 1)[0]
        self.assertIn('Funding Reconciliation', overview)
        self.assertIn('Execution Against Plan', overview)
        self.assertNotIn('<table', overview)

    def test_budget_code_header_matches_green_card_without_farm_details(self):
        from expenses.models import Category
        category = Category.objects.create(name='Seed Purchases', code='SEED')
        self.farm.location = 'Header location'
        self.farm.phone = '+256 700 123456'
        self.farm.email = 'farm@example.com'
        self.farm.save()
        url = reverse('budget_code_detail', kwargs={'pk': category.pk})
        response = self.client.get(url)
        self.assertContains(response, 'budget-toolbar')
        self.assertContains(response, 'Back to Budget Codes')
        self.assertContains(response, 'Print')
        header = response.content.decode().split('<header class="budget-toolbar"', 1)[1].split('</header>', 1)[0]
        for value in [category.name, category.code]:
            self.assertIn(value, header)
        for value in [self.farm.farm_name, self.farm.location, self.farm.phone, self.farm.email]:
            self.assertNotIn(value, header)
        self.assertNotIn('Edit Code', header)
        admin = User.objects.create_superuser(username='code-header-admin', password='test-password')
        self.client.force_login(admin)
        response = self.client.get(url)
        self.assertContains(response, 'Edit Code')
        self.assertContains(response, reverse('model_update', kwargs={
            'app': 'expenses', 'model': 'category', 'pk': category.pk,
        }))

    def test_budget_detail_links_to_registered_edit_form(self):
        edit_url = reverse('model_update', kwargs={
            'app': 'budgets', 'model': 'budget', 'pk': self.budget.pk,
        })
        response = self.client.get(reverse('budget_detail', kwargs={'pk': self.budget.pk}))
        self.assertContains(response, 'Edit Budget')
        self.assertContains(response, reverse('budget_edit', kwargs={'pk': self.budget.pk}))
        self.assertEqual(self.client.get(edit_url).status_code, 200)


class StructuredBudgetTests(TestCase):
    def setUp(self):
        from django.apps import apps
        from expenses.models import BudgetSubCode, Category, ScheduleLine
        self.income_model = apps.get_model('budgets', 'BudgetIncomePlan')
        self.user = User.objects.create_user(username='planner', password='test-password')
        self.farm = FarmProfile.objects.create(user=self.user, farm_name='Planning Farm')
        self.category = Category.objects.create(name='Crop maintenance', code='CROP')
        self.other_category = Category.objects.create(name='Livestock', code='LIV')
        self.subcode = BudgetSubCode.objects.create(category=self.category, code='CM-1', name='Crop Maintenance')
        self.other_subcode = BudgetSubCode.objects.create(category=self.other_category, code='LV-1', name='Livestock Feeding')
        self.bare_subcode = BudgetSubCode.objects.create(category=self.category, code='CM-9', name='Bare head')
        self.schedule_line = ScheduleLine.objects.create(code='OWN-FUNDS', name='Own funds')
        self.item_line = ScheduleLine.objects.create(code='WEED', name='Weed control', subcode=self.subcode)
        self.item_line_b = ScheduleLine.objects.create(code='PEST', name='Pest control', subcode=self.subcode)
        self.other_item_line = ScheduleLine.objects.create(code='FEED', name='Feed mixing', subcode=self.other_subcode)
        self.client.force_login(self.user)

    def payload(self):
        return {
            'farm_profile': self.farm.pk, 'name': 'Operating Budget 2026',
            'start_date': '2026-01-01', 'end_date': '2026-12-31', 'status': 'DRAFT', 'notes': '',
            'incomes-TOTAL_FORMS': '1', 'incomes-INITIAL_FORMS': '0',
            'incomes-0-line_code': 'INC-01', 'incomes-0-funding_type': 'CAPITAL',
            'incomes-0-schedule_line': self.schedule_line.pk, 'incomes-0-planned_amount': '1000',
            'incomes-0-figure_type': 'RECORDED', 'incomes-0-notes': 'Confirmed contribution',
            'items-TOTAL_FORMS': '2', 'items-INITIAL_FORMS': '0',
            'items-0-line_code': 'CROP-01-a', 'items-0-cost_centre': 'Crops',
            'items-0-subcode': self.subcode.pk, 'items-0-reference': 'a',
            'items-0-schedule_line': self.item_line.pk, 'items-0-category': self.category.pk,
            'items-0-planned_amount': '300', 'items-0-figure_type': 'RECORDED',
            'items-0-crop_season': '', 'items-0-notes': 'Quoted cost',
            'items-1-line_code': 'CROP-01-b', 'items-1-cost_centre': 'Crops',
            'items-1-subcode': self.subcode.pk, 'items-1-reference': 'b',
            'items-1-schedule_line': self.item_line_b.pk, 'items-1-category': self.category.pk,
            'items-1-planned_amount': '', 'items-1-figure_type': 'ESTIMATE',
            'items-1-crop_season': '', 'items-1-notes': 'To confirm',
        }

    def test_one_page_create_and_grouped_exports(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from income.models import Income
        data = self.payload()
        data.update({'incomes-TOTAL_FORMS': '2', 'incomes-1-line_code': 'INC-02',
                     'incomes-1-funding_type': 'REVENUE',
                     'incomes-1-schedule_line': self.schedule_line.pk,
                     'incomes-1-planned_amount': '', 'incomes-1-figure_type': 'ESTIMATE',
                     'incomes-1-notes': ''})
        response = self.client.post(reverse('budget_new'), data)
        budget = Budget.objects.get(name='Operating Budget 2026')
        self.assertRedirects(response, reverse('budget_detail', kwargs={'pk': budget.pk}))
        self.assertEqual(budget.items.count(), 2)
        self.assertEqual(budget.total_income_planned, 1000)
        self.assertEqual(budget.total_planned, 300)
        self.assertEqual(budget.funding_balance, 700)
        self.assertEqual(budget.items.get(line_code='CROP-01-b').figure_type, 'ESTIMATE')
        income = budget.income_plans.get(planned_amount=1000)
        self.assertEqual(income.schedule_line, self.schedule_line)
        self.assertEqual(income.name, 'Own funds')
        self.assertEqual(Income.objects.count(), 0)
        page = self.client.get(reverse('budget_detail', kwargs={'pk': budget.pk}))
        for label in ['Sources of Income', 'Crop Maintenance', 'Weed control', 'Subtotal', 'Unconfirmed']:
            self.assertContains(page, label)
        for fmt in ['csv', 'xlsx', 'pdf']:
            export = self.client.get(reverse('budget_detail', kwargs={'pk': budget.pk}), {'export': fmt})
            self.assertEqual(export.status_code, 200)
            self.assertIn('attachment;', export['Content-Disposition'])
            if fmt == 'csv':
                self.assertIn('Planning Farm', export.content.decode())
                self.assertIn('Subtotal - Crops', export.content.decode())
            if fmt == 'xlsx':
                ws = load_workbook(BytesIO(export.content)).active
                self.assertEqual(ws['A1'].value, 'PLANNING FARM')
                self.assertTrue(any(c.value == 'CROP-01-a' for row in ws for c in row))
                recorded = next(row for row in ws if row[0].value == 'CROP-01-a')
                pending = next(row for row in ws if row[0].value == 'CROP-01-b')
                self.assertEqual(recorded[6].font.color.rgb, '001D4ED8')
                self.assertIsNone(pending[6].value)
                self.assertTrue(any(c.data_type == 'f' and c.value.startswith('=SUM(') for row in ws for c in row))

    def test_subcode_syncs_head_and_edit_keeps_it(self):
        self.client.post(reverse('budget_new'), self.payload())
        item = Budget.objects.get(farm_profile=self.farm).items.get(line_code='CROP-01-a')
        self.assertEqual(item.subcode, self.subcode)
        self.assertEqual(item.head, 'Crop Maintenance')
        self.assertEqual(item.schedule_line, self.item_line)
        self.assertEqual(item.line_name, 'Weed control')
        page = self.client.get(reverse('budget_edit', kwargs={'pk': item.budget_id}))
        self.assertContains(page, 'SUBCODE_CATS')
        data = self.payload()
        data.update({'items-INITIAL_FORMS': '2', 'items-0-id': item.pk,
                     'items-1-id': item.budget.items.exclude(pk=item.pk).get().pk,
                     'items-0-subcode': self.subcode.pk,
                     'incomes-INITIAL_FORMS': '1',
                     'incomes-0-id': item.budget.income_plans.get().pk})
        self.assertEqual(self.client.post(reverse('budget_edit', kwargs={'pk': item.budget_id}), data).status_code, 302)
        item.refresh_from_db()
        self.assertEqual(item.subcode, self.subcode)
        self.assertEqual(item.head, 'Crop Maintenance')

    def test_subcode_from_another_budget_code_rejected(self):
        data = self.payload()
        data['items-0-subcode'] = self.other_subcode.pk
        response = self.client.post(reverse('budget_new'), data)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Budget.objects.exists())
        self.assertIn('must belong to the selected budget code',
                      response.context['item_formset'].errors[0]['subcode'][0])

    def test_schedule_line_linked_to_another_subcode_rejected(self):
        data = self.payload()
        data['items-0-schedule_line'] = self.other_item_line.pk
        response = self.client.post(reverse('budget_new'), data)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Budget.objects.exists())
        self.assertIn('linked to a different subcode',
                      response.context['item_formset'].errors[0]['schedule_line'][0])

    def test_multiple_line_items_can_share_one_subcode(self):
        self.client.post(reverse('budget_new'), self.payload())
        budget = Budget.objects.get(farm_profile=self.farm)
        lines = budget.items.filter(subcode=self.subcode)
        self.assertEqual(lines.count(), 2)
        self.assertEqual(set(lines.values_list('schedule_line', flat=True)),
                         {self.item_line.pk, self.item_line_b.pk})

    def test_subcode_row_without_linked_lines_saves_with_subcode_name(self):
        data = self.payload()
        data['items-0-subcode'] = self.bare_subcode.pk
        data['items-0-schedule_line'] = ''
        response = self.client.post(reverse('budget_new'), data)
        budget = Budget.objects.get(name='Operating Budget 2026')
        self.assertRedirects(response, reverse('budget_detail', kwargs={'pk': budget.pk}))
        item = budget.items.get(line_code='CROP-01-a')
        self.assertEqual(item.subcode, self.bare_subcode)
        self.assertEqual(item.line_name, 'Bare head')

    def test_registering_subcode_shows_all_linked_lines_automatically(self):
        data = self.payload()
        data.update({'items-TOTAL_FORMS': '1', 'items-0-subcode': self.subcode.pk,
                     'items-0-schedule_line': ''})
        self.client.post(reverse('budget_new'), data)
        budget = Budget.objects.get(farm_profile=self.farm)
        self.assertEqual(budget.items.count(), 1)
        register = budget.items.get()
        self.assertIsNone(register.schedule_line)
        self.assertEqual(register.subcode, self.subcode)
        self.assertEqual(register.planned_amount, 300)
        from budgets.queries import budget_plan_data
        head = budget_plan_data(budget)['centres'][0]['heads'][0]
        self.assertEqual(head['register'].pk, register.pk)
        self.assertEqual(head['planned'], 300)
        self.assertEqual({line.line_name for line in head['lines']},
                         {'Weed control', 'Pest control'})
        self.assertTrue(all(line.pk is None and line.planned_amount == 0
                            for line in head['lines']))
        page = self.client.get(reverse('budget_detail', kwargs={'pk': budget.pk}))
        for label in ['Subcode total', 'Weed control', 'Pest control']:
            self.assertContains(page, label)

    def test_row_without_subcode_reports_error_on_subcode_field(self):
        data = self.payload()
        data['items-0-subcode'] = ''
        data['items-0-schedule_line'] = ''
        response = self.client.post(reverse('budget_new'), data)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Budget.objects.exists())
        self.assertIn('Pick a subcode', response.context['item_formset'].errors[0]['subcode'][0])

    def test_planner_exposes_subcode_line_map_for_auto_include(self):
        page = self.client.get(reverse('budget_new'))
        self.assertContains(page, 'id="sub-lines"')
        self.assertContains(page, f'"{self.subcode.pk}"')
        self.assertContains(page, 'renderLinePreview')
        self.assertContains(page, '.field-schedule_line { display:none')
        self.assertContains(page, '.budget-line.is-linked { display:none')

    def test_subcode_dropdown_is_limited_to_existing_subcodes(self):
        response = self.client.get(reverse('budget_new'))
        form = response.context['item_formset'].forms[0]
        self.assertEqual(form['subcode'].field.queryset.count(), 3)
        page = self.client.get(reverse('budget_new'))
        self.assertContains(page, 'Select subcode')
        self.assertContains(page, 'SUBCODE_CATS')

    def test_invalid_line_rolls_back_entire_budget(self):
        data = self.payload()
        data['items-0-planned_amount'] = '-1'
        response = self.client.post(reverse('budget_new'), data)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Budget.objects.exists())
        self.assertFalse(self.income_model.objects.exists())

    def test_duplicate_line_codes_rejected(self):
        data = self.payload()
        data['items-1-line_code'] = data['items-0-line_code']
        response = self.client.post(reverse('budget_new'), data)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Budget.objects.exists())

    def test_cross_farm_budget_and_season_rejected(self):
        from crops.models import Crop, CropSeason
        from fields.models import Field
        other = User.objects.create_user(username='other-planner')
        farm = FarmProfile.objects.create(user=other, farm_name='Private Farm')
        budget = Budget.objects.create(farm_profile=farm, name='Private budget')
        self.assertEqual(self.client.get(reverse('budget_edit', kwargs={'pk': budget.pk})).status_code, 404)
        season = CropSeason.objects.create(
            field=Field.objects.create(farm_profile=farm, name='Private field', acreage=1),
            crop=Crop.objects.create(name='Private crop'), planting_date='2026-01-01')
        data = self.payload()
        data['items-0-crop_season'] = season.pk
        self.assertEqual(self.client.post(reverse('budget_new'), data).status_code, 200)
        self.assertFalse(Budget.objects.filter(farm_profile=self.farm).exists())

    def test_existing_report_and_budget_list_exports_still_work(self):
        self.client.post(reverse('budget_new'), self.payload())
        paths = ['/reports/', '/reports/sales/', '/reports/production/',
                 '/reports/livestock/', '/reports/operations/',
                 '/dashboard/budgets/budget/', '/dashboard/budgets/budgetitem/']
        for path in paths:
            for fmt in ['csv', 'xlsx', 'pdf']:
                with self.subTest(path=path, format=fmt):
                    response = self.client.get(path, {'export': fmt})
                    self.assertEqual(response.status_code, 200)
                    self.assertIn('attachment;', response['Content-Disposition'])

    def test_reversed_dates_do_not_save_the_plan(self):
        data = self.payload()
        data['end_date'] = '2025-01-01'
        response = self.client.post(reverse('budget_new'), data)
        self.assertContains(response, 'End date cannot be before the start date')
        self.assertFalse(Budget.objects.exists())

    def test_generated_codes_do_not_collide_with_manual_codes(self):
        data = self.payload()
        data['items-0-line_code'] = ''
        data['items-1-line_code'] = 'EXP-001'
        self.assertEqual(self.client.post(reverse('budget_new'), data).status_code, 302)
        codes = set(Budget.objects.get(farm_profile=self.farm).items.values_list('line_code', flat=True))
        self.assertEqual(codes, {'EXP-001', 'EXP-002'})

    def test_edit_updates_lines_without_duplicates(self):
        self.client.post(reverse('budget_new'), self.payload())
        budget = Budget.objects.get(farm_profile=self.farm)
        first, second = list(budget.items.order_by('pk'))
        income = budget.income_plans.get()
        data = self.payload()
        data.update({'items-INITIAL_FORMS': '2', 'items-0-id': first.pk, 'items-1-id': second.pk,
                     'incomes-INITIAL_FORMS': '1', 'incomes-0-id': income.pk,
                     'items-1-planned_amount': '200', 'incomes-0-planned_amount': '1500', 'status': 'ACTIVE'})
        response = self.client.post(reverse('budget_edit', kwargs={'pk': budget.pk}), data)
        self.assertEqual(response.status_code, 302)
        budget.refresh_from_db()
        self.assertEqual(budget.items.count(), 2)
        self.assertEqual(budget.income_plans.count(), 1)
        self.assertEqual(budget.total_planned, 500)
        self.assertEqual(budget.funding_balance, 1000)
        self.assertEqual(budget.status, 'ACTIVE')

    def test_unconfirmed_budget_cannot_be_activated(self):
        data = self.payload()
        data['status'] = 'ACTIVE'
        data['items-1-schedule_line'] = ''
        response = self.client.post(reverse('budget_new'), data)
        self.assertContains(response, 'Confirm all line amounts before activating')
        self.assertFalse(Budget.objects.exists())

    def test_posted_line_cannot_be_removed_and_edits_are_atomic(self):
        from datetime import date
        from expenses.models import Expense
        self.client.post(reverse('budget_new'), self.payload())
        budget = Budget.objects.get(farm_profile=self.farm)
        first, second = list(budget.items.order_by('pk'))
        income = budget.income_plans.get()
        Expense.objects.create(farm_profile=self.farm, category=self.category,
                               budget_item=first, amount=50, date=date(2026, 5, 1))
        data = self.payload()
        data.update({'name': 'Should not save', 'items-INITIAL_FORMS': '2',
                     'items-0-id': first.pk, 'items-1-id': second.pk,
                     'items-0-DELETE': 'on', 'incomes-INITIAL_FORMS': '1', 'incomes-0-id': income.pk})
        response = self.client.post(reverse('budget_edit', kwargs={'pk': budget.pk}), data)
        self.assertContains(response, 'Lines with posted expenses cannot be removed')
        budget.refresh_from_db()
        self.assertEqual(budget.name, 'Operating Budget 2026')
        self.assertEqual(budget.items.count(), 2)
        self.assertEqual(budget.total_spent, 50)
        self.assertEqual(budget.remaining, 250)

    def test_editor_cannot_adopt_a_line_from_another_budget(self):
        self.client.post(reverse('budget_new'), self.payload())
        original = Budget.objects.get(farm_profile=self.farm)
        data = self.payload()
        data['name'] = 'Forged budget'
        data.update({'items-INITIAL_FORMS': '1', 'items-0-id': original.items.first().pk})
        response = self.client.post(reverse('budget_new'), data)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Budget.objects.filter(name='Forged budget').exists())

    def test_formula_like_text_is_not_executed_in_exports(self):
        from io import BytesIO
        from openpyxl import load_workbook
        data = self.payload()
        data['name'] = 'Plan & Farm <2026>'
        data['items-0-reference'] = '=1+1'
        response = self.client.post(reverse('budget_new'), data)
        self.assertEqual(response.status_code, 302)
        budget = Budget.objects.get(farm_profile=self.farm)
        url = reverse('budget_detail', kwargs={'pk': budget.pk})
        ws = load_workbook(BytesIO(self.client.get(url, {'export': 'xlsx'}).content)).active
        self.assertEqual(next(c for row in ws for c in row if c.value == '=1+1').data_type, 's')
        self.assertIn("'=1+1", self.client.get(url, {'export': 'csv'}).content.decode())
        self.assertEqual(self.client.get(url, {'export': 'pdf'}).status_code, 200)

    def test_worker_cannot_create_budget(self):
        from accounts.models import FarmMembership
        worker = User.objects.create_user(username='planning-worker')
        FarmMembership.objects.create(user=worker, farm_profile=self.farm, role='WORKER')
        self.client.force_login(worker)
        self.assertEqual(self.client.post(reverse('budget_new'), self.payload()).status_code, 403)


class BudgetMigrationTests(TransactionTestCase):
    def test_legacy_amounts_funds_and_expense_links_survive_migration(self):
        from datetime import date
        from django.db import connection
        from django.db.migrations.executor import MigrationExecutor
        from expenses.models import Category, Expense
        user = User.objects.create_user(username='legacy-planner')
        farm = FarmProfile.objects.create(user=user, farm_name='Legacy Farm')
        category = Category.objects.create(name='Legacy seeds', code='LEGACY')
        executor = MigrationExecutor(connection)
        latest = executor.loader.graph.leaf_nodes('budgets')
        # Rolling budgets back to 0004 also unapplies dependent expenses
        # migrations, so bring every app forward again when restoring.
        all_leaf = executor.loader.graph.leaf_nodes()
        previous = [('budgets', '0004_budgetitem_fund_amount')]
        try:
            executor.migrate(previous)
            old_apps = executor.loader.project_state(previous).apps
            OldBudget = old_apps.get_model('budgets', 'Budget')
            OldItem = old_apps.get_model('budgets', 'BudgetItem')
            budget = OldBudget.objects.create(name='Legacy plan', farm_profile_id=farm.pk)
            item = OldItem.objects.create(budget_id=budget.pk, category_id=category.pk,
                                          planned_amount=500, fund_amount=800, notes='Historical basis')
            executor = MigrationExecutor(connection)
            executor.migrate(all_leaf)
            migrated = Budget.objects.get(pk=budget.pk)
            line = migrated.items.get(pk=item.pk)
            expense = Expense.objects.create(farm_profile=farm, category=category,
                                             budget_item_id=line.pk, amount=125, date=date(2026, 1, 1))
            self.assertEqual(line.planned_amount, 500)
            self.assertEqual(line.fund_amount, 800)
            self.assertEqual(line.notes, 'Historical basis')
            self.assertEqual(line.line_name, category.name)
            self.assertEqual(line.figure_type, 'RECORDED')
            self.assertTrue(line.line_code.startswith('LEG-'))
            self.assertEqual(expense.budget_item_id, line.pk)
            self.assertEqual(migrated.total_spent, 125)
            self.assertEqual(migrated.remaining, 375)
            self.assertEqual(migrated.income_plans.count(), 0)
        finally:
            MigrationExecutor(connection).migrate(all_leaf)



class BudgetBurnChartTests(TestCase):
    def setUp(self):
        from datetime import date
        from decimal import Decimal
        from expenses.models import Category, Expense
        from budgets.models import BudgetItem
        self.user = User.objects.create_user(username='burn-owner', password='pw')
        self.farm = FarmProfile.objects.create(user=self.user, farm_name='Burn Farm')
        self.budget = Budget.objects.create(
            farm_profile=self.farm, name='FY26 Burn',
            start_date=date(2026, 1, 1), end_date=date(2026, 4, 30))
        category = Category.objects.create(name='Burn cat')
        item = BudgetItem.objects.create(
            budget=self.budget, category=category, line_code='B-1',
            planned_amount=Decimal('1200'))
        Expense.objects.create(
            farm_profile=self.farm, budget_item=item, category=category,
            date=date(2026, 1, 15), amount=Decimal('300'))
        Expense.objects.create(
            farm_profile=self.farm, budget_item=item, category=category,
            date=date(2026, 3, 10), amount=Decimal('600'))
        self.client.force_login(self.user)

    def test_overview_shows_burn_chart(self):
        from decimal import Decimal
        response = self.client.get(
            reverse('budget_detail', kwargs={'pk': self.budget.pk}))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '<article class="burn-chart">')
        self.assertContains(response, 'Spend vs Plan')
        chart = response.context['burn_chart']
        self.assertEqual(len(chart['rows']), 4)
        self.assertEqual(chart['rows'][-1]['actual'], Decimal('900'))
        self.assertEqual(chart['rows'][-1]['plan'], Decimal('1200'))
        self.assertTrue(chart['has_actual'])
        alloc = response.context['alloc_segments']
        self.assertEqual(len(alloc), 1)
        self.assertEqual(alloc[0]['share'], 100.0)
        self.assertContains(response, 'Budget Allocation')

    def test_chart_hidden_without_dates_or_expenses(self):
        bare = Budget.objects.create(farm_profile=self.farm, name='Bare')
        response = self.client.get(
            reverse('budget_detail', kwargs={'pk': bare.pk}))
        self.assertEqual(response.status_code, 200)
        self.assertIsNone(response.context['burn_chart'])
        self.assertNotContains(response, '<article class="burn-chart">')
        self.assertNotContains(response, 'class="alloc-chart"')
